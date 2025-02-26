#Opens a template workbook and creates a copy of a template worksheet
#within a copy of the workbook

#Import libraries
import os
from openpyxl import load_workbook
import openpyxl
import shutil
from datetime import datetime
import json
import csv
#from PIL import Image
#import PIL

#row number :
r = 21

def convertTFtoint(stringin) -> int:
    output = 0
    if stringin == "TRUE":
        output = 1
    return output

#Print Script Start
print ('!!!!!!!!! Script Start !!!!!!!!!')

#Changes Working Directory To a Known Location
#os.chdir('C:\Users\es2433\Desktop\PV_PythonSandbox\ExcelBookLayout')

#Copy Excel file from template file with datetime stamp
CurrDateTime = str(datetime.now())
CurrDateTime = CurrDateTime.replace(':','-')
CurrDateTime = CurrDateTime.replace(' ','_')
CurrDateTime = CurrDateTime.replace('.','-')
shutil.copy('ScoutingData_template.xlsx','ScoutingData_'+CurrDateTime+'.xlsx')

#Load Match CSV To Match Database
cwd = os.getcwd()
files = os.listdir(cwd)
newest_file_match = None
newest_time_match = 0
newest_file_pit = None
newest_time_pit = 0

for file in files:
    if 'data@' in file:
        file_path = os.path.join(cwd,file)
        creation_time = os.path.getctime(file_path)

        if creation_time > newest_time_match:
            newest_time_match = creation_time
            newest_file_match = file_path

    if 'data_pit@' in file:
        file_path = os.path.join(cwd,file)
        creation_time = os.path.getctime(file_path)

        if creation_time > newest_time_pit:
            newest_time_pit = creation_time
            newest_file_pit = file_path


matchcsv = newest_file_match
pitcsv = newest_file_pit

#Opens Excel Spreadsheet
wb = load_workbook('ScoutingData_'+CurrDateTime+'.xlsx')

#Find template worksheet by name
ws_template = wb['template']#.get_sheet_by_name(name = 'template')

#Read in team numbers from team list
wsteam_list = wb['team_list']#.get_sheet_by_name(name = 'team_list')
TeamList=[]
TeamNameList=[]

#Populate Team List Based Upon Excel Data
for rowindex in range(wsteam_list.max_row):
    TeamList.append(str(wsteam_list.cell(row = rowindex + 1,column = 1).value))
    TeamNameList.append(str(wsteam_list.cell(row = rowindex + 1,column = 2).value))


#Loop through team list creating worksheet for each one
for inc, team in enumerate(TeamList):
#Create copy of template sheet
    ws2 = wb.copy_worksheet(ws_template)

#Rename Newly Copied Sheet
    ws2.title = team
#----------------------------------------

    #Fill out data on sheet
    ws2.cell(row=1,column=2).value = team
    ws2.cell(row=2,column=2).value = TeamNameList[inc]
    
    #Put in pit scouting data
    rows_pit = []
    with open(pitcsv, 'r',newline="") as csvfile:
        reader = csv.DictReader(csvfile)
        for row in reader:
            if row["Team Number"][:-2] == team:
                rows_pit.append(row)

    for inc3, row_pit in enumerate(rows_pit):

        #Generate Coral String
        coral_string = ""
        if row_pit['Coral L1'] == 'True':
            coral_string += "L1"+","
        if row_pit['Coral L2'] == 'True':
            coral_string += "L2"+","
        if row_pit['Coral L3'] == 'True':
            coral_string += "L3"+","
        if row_pit['Coral L4'] == 'True':
            coral_string += "L4"+","
        if row_pit['Coral HP Pickup'] == 'True':
            coral_string += "HP Pickup"+","
        if row_pit[' Coral Floor Pickup'] == 'True':
            coral_string += "Floor Pickup"+","
        if coral_string != "":
            coral_string = coral_string[:-1]
        else:
            coral_string = "None Specified"
        
        #Generate Algae String
        algae_string = ""
        if row_pit['Algae Processor'] == 'True':
            algae_string += "Processor"+","
        if row_pit['Algae Barge'] == 'True':
            algae_string += "Barge"+","
        if row_pit['Algae Displaced'] == 'True':
            algae_string += "Displace from Reef"+","
        if row_pit['Algae Ground Pickup'] == 'True':
            algae_string += "Floor Pickup"+","
        if algae_string != "":
            algae_string = algae_string[:-1]
        else:
            algae_string = "None Specified"
        
        auto_string = ""
        if row_pit['Auto Leave'] == 'True':
            auto_string += "Auto Leave"+","
        auto_string += row_pit['Auto Info'].replace("$n", "")
        if auto_string == "":
            auto_string = "None Specified"

        cage_string = ""
        if row_pit['Shallow Cage'] == 'True':
            cage_string += "Shallow Cage"+","
        if row_pit['Deep Cage'] == 'True':
            cage_string += "Deep Cage"+","
        if cage_string != "":
            cage_string = cage_string[:-1]
        else:
            cage_string = "None Specified"

        ws2.cell(row=inc3+5,column=2).value = row_pit['Drivetrain'].replace("$n", "")
        ws2.cell(row=inc3+6,column=2).value = coral_string
        ws2.cell(row=inc3+7,column=2).value = algae_string
        ws2.cell(row=inc3+8,column=2).value = auto_string
        ws2.cell(row=inc3+9,column=2).value = cage_string
        ws2.cell(row=inc3+10,column=2).value = row_pit[' Fit Under Shallow']
        ws2.cell(row=inc3+11,column=2).value = row_pit['Best Aspect'].replace("$n", "")
        ws2.cell(row=inc3+12,column=2).value = row_pit['Comments'].replace("$n", "")

    #Put in match scouting data
    rows = []
    with open(matchcsv, 'r',newline="") as csvfile:
        reader = csv.DictReader(csvfile)
        for row in reader:
            if row["Team Number"][:-2] == team:
                rows.append(row)

    for inc2, row in enumerate(rows):
    							 	 		
        ws2.cell(row=inc2+r,column=1).value = int(row['Match Number'])
        ws2.cell(row=inc2+r,column=2).value = row['Alliance']
        ws2.cell(row=inc2+r,column=3).value = row['Name'].replace("$n", "")
        ws2.cell(row=inc2+r,column=4).value = int(row['Auto Coral L1'])
        ws2.cell(row=inc2+r,column=5).value = int(row['Auto Coral L2'])
        ws2.cell(row=inc2+r,column=6).value = int(row['Auto Coral L3'])
        ws2.cell(row=inc2+r,column=7).value = int(row['Auto Coral L4'])
        ws2.cell(row=inc2+r,column=8).value = int(row['Auto Algae Displaced'])
        ws2.cell(row=inc2+r,column=9).value = int(row['Auto Algae Processor'])
        ws2.cell(row=inc2+r,column=10).value = int(row['Auto Algae Barge'])
        ws2.cell(row=inc2+r,column=11).value = convertTFtoint(row['Auto Left'])
        ws2.cell(row=inc2+r,column=12).value = int(row['Teleop Coral L1'])
        ws2.cell(row=inc2+r,column=13).value = int(row['Teleop Coral L2'])
        ws2.cell(row=inc2+r,column=14).value = int(row['Teleop Coral L3'])
        ws2.cell(row=inc2+r,column=15).value = int(row['Teleop Coral L4'])
        ws2.cell(row=inc2+r,column=16).value = int(row['Teleop Algae Displaced'])
        ws2.cell(row=inc2+r,column=17).value = int(row['Teleop Algae Processed'])
        ws2.cell(row=inc2+r,column=18).value = int(row['Teleop Algae Barge'])
        ws2.cell(row=inc2+r,column=19).value = row['Defense']
        ws2.cell(row=inc2+r,column=20).value = row['Penalties'].replace("$n", "")
        ws2.cell(row=inc2+r,column=21).value = convertTFtoint(row['Attempt Shallow'])
        ws2.cell(row=inc2+r,column=22).value = convertTFtoint(row['Sucess Shallow'])
        ws2.cell(row=inc2+r,column=23).value = convertTFtoint(row[' Attempt Deep'])
        ws2.cell(row=inc2+r,column=24).value = convertTFtoint(row[' Success Deep'])
        ws2.cell(row=inc2+r,column=25).value = convertTFtoint(row['Parked'])
        ws2.cell(row=inc2+r,column=26).value = row['Comments'].replace("$n", "")


    # Add in Summary Page Info
    ws_sum = wb['Summary']#.get_sheet_by_name(name = 'Summary')

    #Adds information to summary sheet from team worksheet
    ws_sum.cell(row = inc+9,column = 1).value = "='"+team+"'!"+"B1" #Team Number

    #AVG
    ws_sum.cell(row = inc+9,column = 2).value = "='"+team+"'!"+"D17" #Autos
    ws_sum.cell(row = inc+9,column = 3).value = "='"+team+"'!"+"E17"
    ws_sum.cell(row = inc+9,column = 4).value = "='"+team+"'!"+"F17"
    ws_sum.cell(row = inc+9,column = 5).value = "='"+team+"'!"+"G17"
    ws_sum.cell(row = inc+9,column = 6).value = "='"+team+"'!"+"H17"
    ws_sum.cell(row = inc+9,column = 7).value = "='"+team+"'!"+"I17"
    ws_sum.cell(row = inc+9,column = 8).value = "='"+team+"'!"+"J17"
    ws_sum.cell(row = inc+9,column = 9).value = "='"+team+"'!"+"K17"
                                                                    #Tele
    ws_sum.cell(row = inc+9,column = 10).value = "='"+team+"'!"+"L17"
    ws_sum.cell(row = inc+9,column = 11).value = "='"+team+"'!"+"M17"
    ws_sum.cell(row = inc+9,column = 12).value = "='"+team+"'!"+"N17"
    ws_sum.cell(row = inc+9,column = 13).value = "='"+team+"'!"+"O17"
    ws_sum.cell(row = inc+9,column = 14).value = "='"+team+"'!"+"P17"
    ws_sum.cell(row = inc+9,column = 15).value = "='"+team+"'!"+"Q17"
    ws_sum.cell(row = inc+9,column = 16).value = "='"+team+"'!"+"R17"

    ws_sum.cell(row = inc+9,column = 17).value = "='"+team+"'!"+"U17" #Endgame
    ws_sum.cell(row = inc+9,column = 18).value = "='"+team+"'!"+"V17"
    ws_sum.cell(row = inc+9,column = 19).value = "='"+team+"'!"+"W17"                                                                
    ws_sum.cell(row = inc+9,column = 20).value = "='"+team+"'!"+"X17"
    ws_sum.cell(row = inc+9,column = 21).value = "='"+team+"'!"+"Y17"

    #MAX
    ws_sum.cell(row = inc+9,column = 22).value = "='"+team+"'!"+"D18" #Autos
    ws_sum.cell(row = inc+9,column = 23).value = "='"+team+"'!"+"E18"
    ws_sum.cell(row = inc+9,column = 24).value = "='"+team+"'!"+"F18"
    ws_sum.cell(row = inc+9,column = 25).value = "='"+team+"'!"+"G18"
    ws_sum.cell(row = inc+9,column = 26).value = "='"+team+"'!"+"H18"
    ws_sum.cell(row = inc+9,column = 27).value = "='"+team+"'!"+"I18"
    ws_sum.cell(row = inc+9,column = 28).value = "='"+team+"'!"+"J18"
    ws_sum.cell(row = inc+9,column = 29).value = "='"+team+"'!"+"K18"
                                                                    #Tele
    ws_sum.cell(row = inc+9,column = 30).value = "='"+team+"'!"+"L18"
    ws_sum.cell(row = inc+9,column = 31).value = "='"+team+"'!"+"M18"
    ws_sum.cell(row = inc+9,column = 32).value = "='"+team+"'!"+"N18"
    ws_sum.cell(row = inc+9,column = 33).value = "='"+team+"'!"+"O18"
    ws_sum.cell(row = inc+9,column = 34).value = "='"+team+"'!"+"P18"
    ws_sum.cell(row = inc+9,column = 35).value = "='"+team+"'!"+"Q18"
    ws_sum.cell(row = inc+9,column = 36).value = "='"+team+"'!"+"R18"

    ws_sum.cell(row = inc+9,column = 37).value = "='"+team+"'!"+"U18" #Endgame
    ws_sum.cell(row = inc+9,column = 38).value = "='"+team+"'!"+"V18"
    ws_sum.cell(row = inc+9,column = 39).value = "='"+team+"'!"+"W18"                                                                
    ws_sum.cell(row = inc+9,column = 40).value = "='"+team+"'!"+"X18"
    ws_sum.cell(row = inc+9,column = 41).value = "='"+team+"'!"+"Y18"

    #MIN
    ws_sum.cell(row = inc+9,column = 42).value = "='"+team+"'!"+"D19" #Autos
    ws_sum.cell(row = inc+9,column = 43).value = "='"+team+"'!"+"E19"
    ws_sum.cell(row = inc+9,column = 44).value = "='"+team+"'!"+"F19"
    ws_sum.cell(row = inc+9,column = 45).value = "='"+team+"'!"+"G19"
    ws_sum.cell(row = inc+9,column = 46).value = "='"+team+"'!"+"H19"
    ws_sum.cell(row = inc+9,column = 47).value = "='"+team+"'!"+"I19"
    ws_sum.cell(row = inc+9,column = 48).value = "='"+team+"'!"+"J19"
    ws_sum.cell(row = inc+9,column = 49).value = "='"+team+"'!"+"K19"
                                                                    #Tele
    ws_sum.cell(row = inc+9,column = 50).value = "='"+team+"'!"+"L19"
    ws_sum.cell(row = inc+9,column = 51).value = "='"+team+"'!"+"M19"
    ws_sum.cell(row = inc+9,column = 52).value = "='"+team+"'!"+"N19"
    ws_sum.cell(row = inc+9,column = 53).value = "='"+team+"'!"+"O19"
    ws_sum.cell(row = inc+9,column = 54).value = "='"+team+"'!"+"P19"
    ws_sum.cell(row = inc+9,column = 55).value = "='"+team+"'!"+"Q19"
    ws_sum.cell(row = inc+9,column = 56).value = "='"+team+"'!"+"R19"

    ws_sum.cell(row = inc+9,column = 57).value = "='"+team+"'!"+"U19" #Endgame
    ws_sum.cell(row = inc+9,column = 58).value = "='"+team+"'!"+"V19"
    ws_sum.cell(row = inc+9,column = 59).value = "='"+team+"'!"+"W19"                                                                
    ws_sum.cell(row = inc+9,column = 60).value = "='"+team+"'!"+"X19"
    ws_sum.cell(row = inc+9,column = 61).value = "='"+team+"'!"+"Y19"

    #Pits
    ws_sum.cell(row = inc+9,column = 62).value = "='"+team+"'!"+"S17" #Defense
    ws_sum.cell(row = inc+9,column = 63).value = "='"+team+"'!"+"B5" #DriveTrain


""" 
###Put in picture ----------------------------------
##    if (os.path.isfile('Pic/'+team+'.jpg')): #Check to make sure picture exists
##        teamnum = team
##    else:
##        teamnum ='default' #Stand in Photo For Teams With Missing Picture
##
##    #print teamnum
##    #anchor = 'G1'
##    
##    im = Image.open('Pic/'+teamnum+'.jpg')
##    image2 = im.resize((250,250),Image.ANTIALIAS)
##    image2.save(teamnum + "_thumbnail.jpg", "JPEG")
##    img = openpyxl.drawing.image.Image(teamnum + "_thumbnail.jpg")
##    ws2.add_image(img,'G3')
# ------------------------------------------------------
"""
#Save workbook
wb.save('ScoutingData_'+CurrDateTime+'.xlsx')

#Print Script End
print ('!!!!!!!!! File Created !!!!!!!!!')
